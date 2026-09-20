"""Parse and upsert clients from CSV / Excel rate sheets."""
from __future__ import annotations

import csv
import io
import re
from decimal import Decimal, InvalidOperation
from typing import Any

from sqlmodel import Session, select

from models.client import Client
from models.enums import ClientContractStatusEnum, ClientOwnerTypeEnum

DEFAULT_PLATFORM = "Unassigned"

# Spreadsheet headers → canonical keys (lowercase, stripped).
_HEADER_ALIASES: dict[str, str] = {
    "client": "name",
    "client name": "name",
    "name": "name",
    "billing rate usd/hr": "billing_rate_usd",
    "billing rate usd": "billing_rate_usd",
    "billing rate": "billing_rate_usd",
    "rate usd/hr": "billing_rate_usd",
    "rate": "billing_rate_usd",
    "billing_rate_usd": "billing_rate_usd",
    "active?": "active",
    "active": "active",
    "contract status": "contract_status",
    "contract_status": "contract_status",
    "status": "contract_status",
    "notes": "notes",
    "note": "notes",
    "platform": "platform",
    "account email": "account_email",
    "account_email": "account_email",
    "email": "account_email",
    "account id": "account_id",
    "account_id": "account_id",
    "login reference": "login_reference",
    "login_reference": "login_reference",
}


def _norm_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def parse_billing_rate(value: Any) -> Decimal | None:
    """Accept `$18.70`, `18.70`, blanks → None."""
    text = _cell_str(value)
    if not text:
        return None
    cleaned = text.replace(",", "").replace("$", "").strip()
    cleaned = re.sub(r"(?i)\s*usd(?:\s*/\s*hr)?$", "", cleaned).strip()
    if not cleaned:
        return None
    try:
        amount = Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid billing rate: {value!r}") from exc
    if amount < 0:
        raise ValueError(f"Billing rate cannot be negative: {value!r}")
    return amount.quantize(Decimal("0.01"))


def parse_active_status(value: Any) -> ClientContractStatusEnum | None:
    text = _cell_str(value)
    if not text:
        return None
    lowered = text.lower()
    if lowered in {"yes", "y", "true", "1", "active"}:
        return ClientContractStatusEnum.active
    if lowered in {"ended", "end", "terminated", "closed"}:
        return ClientContractStatusEnum.ended
    # Anything else (No, paused, inactive, …) → paused
    return ClientContractStatusEnum.paused


def _map_headers(raw_headers: list[Any]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for idx, header in enumerate(raw_headers):
        key = _HEADER_ALIASES.get(_norm_header(header))
        if key:
            mapping[idx] = key
    return mapping


def _row_dict(cells: list[Any], col_map: dict[int, str]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for idx, key in col_map.items():
        if idx < len(cells):
            out[key] = cells[idx]
    return out


def rows_from_csv(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    try:
        headers = next(reader)
    except StopIteration:
        return []
    col_map = _map_headers(headers)
    if "name" not in col_map.values():
        raise ValueError('Spreadsheet must include a "Client" (or Name) column.')
    rows: list[dict[str, Any]] = []
    for cells in reader:
        if not any(_cell_str(c) for c in cells):
            continue
        rows.append(_row_dict(cells, col_map))
    return rows


def rows_from_xlsx(data: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Excel import requires openpyxl on the server.") from exc

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        try:
            headers = list(next(iterator))
        except StopIteration:
            return []
        col_map = _map_headers(headers)
        if "name" not in col_map.values():
            raise ValueError('Spreadsheet must include a "Client" (or Name) column.')
        rows: list[dict[str, Any]] = []
        for cells in iterator:
            values = list(cells or ())
            if not any(_cell_str(c) for c in values):
                continue
            rows.append(_row_dict(values, col_map))
        return rows
    finally:
        wb.close()


def parse_client_import_file(filename: str, data: bytes) -> list[dict[str, Any]]:
    lower = (filename or "").lower()
    if lower.endswith(".csv"):
        return rows_from_csv(data)
    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        return rows_from_xlsx(data)
    if lower.endswith(".xls"):
        raise ValueError("Legacy .xls is not supported — save as .xlsx or CSV.")
    # Sniff: ZIP/xlsx starts with PK; otherwise treat as CSV.
    if data[:2] == b"PK":
        return rows_from_xlsx(data)
    return rows_from_csv(data)


def upsert_clients_from_rows(db: Session, rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Create or update clients by name (case-insensitive). Empty cells do not clear existing values."""
    existing = db.exec(select(Client)).all()
    by_name = {c.name.strip().lower(): c for c in existing if c.name}

    created = 0
    updated = 0
    skipped = 0
    errors: list[str] = []

    for index, raw in enumerate(rows, start=2):  # row 1 = header in sheets
        try:
            name = _cell_str(raw.get("name"))
            if not name:
                skipped += 1
                continue

            rate = parse_billing_rate(raw.get("billing_rate_usd"))
            status = parse_active_status(raw.get("active"))
            if status is None and "contract_status" in raw:
                status = parse_active_status(raw.get("contract_status"))

            platform = _cell_str(raw.get("platform"))
            notes = _cell_str(raw.get("notes")) or None
            account_email = _cell_str(raw.get("account_email")) or None
            account_id = _cell_str(raw.get("account_id")) or None
            login_reference = _cell_str(raw.get("login_reference")) or None

            key = name.lower()
            client = by_name.get(key)
            if client is None:
                client = Client(
                    name=name,
                    platform=platform or DEFAULT_PLATFORM,
                    billing_rate_usd=rate,
                    account_email=account_email,
                    account_id=account_id,
                    login_reference=login_reference,
                    owner_type=ClientOwnerTypeEnum.gs,
                    contract_status=status or ClientContractStatusEnum.active,
                    notes=notes,
                )
                db.add(client)
                by_name[key] = client
                created += 1
            else:
                # Update only fields that have values in this row.
                client.name = name
                if platform:
                    client.platform = platform
                if rate is not None:
                    client.billing_rate_usd = rate
                if status is not None:
                    client.contract_status = status
                if notes is not None:
                    client.notes = notes
                if account_email is not None:
                    client.account_email = account_email
                if account_id is not None:
                    client.account_id = account_id
                if login_reference is not None:
                    client.login_reference = login_reference
                db.add(client)
                updated += 1
        except ValueError as exc:
            errors.append(f"Row {index}: {exc}")
            skipped += 1

    db.commit()
    return {
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors[:50],
        "total_rows": len(rows),
    }
